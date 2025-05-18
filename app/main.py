import streamlit as st
import os
import pandas as pd
import docx
import openpyxl
from io import BytesIO
from dotenv import load_dotenv
from azure.storage.blob import BlobServiceClient
from ocr import extract_text_from_pdf
from embedder import process_text_to_faiss, query_faiss_index
from pathlib import Path

# Load environment variables
env_path = Path(__file__).resolve().parent.parent / '.env'
load_dotenv(dotenv_path=env_path)
conn_str = os.getenv("AZURE_STORAGE_CONNECTION_STRING")
blob_service_client = BlobServiceClient.from_connection_string(conn_str)
container_name = "raw-documents"

def extract_text_from_docx(file):
    doc = docx.Document(file)
    return "\n".join([para.text for para in doc.paragraphs])

def upload_to_azure(file, filename):
    blob_client = blob_service_client.get_blob_client(container=container_name, blob=filename)
    blob_client.upload_blob(file, overwrite=True)
    st.success(f"✅ Uploaded to Azure Blob Storage: `{filename}`")

def main():
    st.title("📂 Upload your Files for SR&ED AI Tool")

    uploaded_files = st.file_uploader(
        "Upload PDFs, CSVs, Word docs, Excel files, or Audio files (mp3, wav)",
        type=["pdf", "csv", "docx", "xlsx", "mp3", "wav"],
        accept_multiple_files=True
    )

    all_extracted_text = ""

    if uploaded_files:
        st.write(f"Uploaded {len(uploaded_files)} file(s):")
        for uploaded_file in uploaded_files:
            st.write(f"- {uploaded_file.name} ({uploaded_file.type}, {uploaded_file.size} bytes)")
            filename = uploaded_file.name.lower()

            # Upload to Azure
            upload_to_azure(uploaded_file, filename)

            if filename.endswith(".csv"):
                df = pd.read_csv(uploaded_file)
                st.dataframe(df.head())

            elif filename.endswith(".pdf"):
                st.write("PDF uploaded — ready for OCR processing.")
                with open("temp.pdf", "wb") as f:
                    f.write(uploaded_file.getbuffer())
                extracted_text = extract_text_from_pdf("temp.pdf")
                st.text_area("Extracted Text", extracted_text, height=300)
                all_extracted_text += extracted_text + "\n\n"

            elif filename.endswith(".docx"):
                st.write("Word document uploaded — extracting text.")
                text = extract_text_from_docx(uploaded_file)
                st.text_area("Extracted Text", text, height=300)
                all_extracted_text += text + "\n\n"

            elif filename.endswith(".xlsx"):
                st.write("Excel file uploaded — previewing first sheet.")
                wb = openpyxl.load_workbook(uploaded_file, read_only=True)
                sheet = wb.active
                data = sheet.values
                columns = next(data)
                df = pd.DataFrame(data, columns=columns)
                st.dataframe(df.head())

            elif filename.endswith(".mp3") or filename.endswith(".wav"):
                st.write("Audio file uploaded — ready for transcription.")
                # Add audio transcription here

            else:
                st.write("Unsupported file type.")

        if all_extracted_text.strip():
            st.write("Embedding all extracted text into FAISS index...")
            process_text_to_faiss(all_extracted_text, index_path="data/faiss_index")

    query = st.text_input("Ask a question:")
    if query:
        try:
            results = query_faiss_index(query)
            for i, (doc, score) in enumerate(results):
                st.markdown(f"**Result {i+1} (Score: {score:.2f})**")
                st.write(doc.page_content)
                st.markdown("---")
        except Exception as e:
            st.error(str(e))


if __name__ == "__main__":
    main()
